"""Build separate January 2026 human-readable DOCX guides for each process.

Markdown under ``docs/`` is the maintained source.
This module assembles reference copies without altering source Markdown.
"""

from __future__ import annotations

import os
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote, unquote, urlsplit

from docx import Document as WordDocument
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "generated"


@dataclass(frozen=True)
class Document:
    stem: str
    title: str
    pages: tuple[Path, ...]


DOCUMENTS = (
    Document(
        stem="PRA-Jan-2026",
        title="PRA Runbook",
        pages=(DOCS / "pra" / "runbook.md",),
    ),
    Document(
        stem="BSCR-Jan-2026",
        title="BSCR Schedule X Runbook",
        pages=(DOCS / "bscr" / "runbook.md",),
    ),
    Document(
        stem="Lloyds-Supplementary-Jan-2026",
        title="Lloyd's/RDS Supplementary Runbook",
        pages=(DOCS / "lloyds" / "runbook.md",),
    ),
    Document(
        stem="Global-Exposures-Jan-2026",
        title="Global Exposures Runbook",
        pages=(DOCS / "global-exposures" / "runbook.md",),
    ),
)


@dataclass(frozen=True)
class PreparedPage:
    path: Path
    text: str
    prefix: str
    page_anchor: str
    anchors: dict[str, tuple[str, ...]]


@dataclass
class RenderState:
    bookmark_names: dict[str, str]
    next_bookmark_id: int = 1


_HEADING = re.compile(r"^(?P<marks>#{1,6})\s+(?P<text>.*?)(?:\s+#+)?\s*$")
_RENDER_HEADING = re.compile(
    r"^(?P<marks>#{1,6})\s+(?P<text>.*?)(?:\s+\{#(?P<anchor>[A-Za-z0-9_-]+)\})?\s*$"
)
_FENCE = re.compile(r"^(?P<indent>\s*)(?P<mark>`{3,}|~{3,})")
_LINK = re.compile(
    r"(?<!!)(?P<full>\[(?P<label>[^\]]+)\]\((?P<target><[^>]*>|[^)\s]+)(?P<tail>[^)]*)\))"
)
_INLINE_MARK = re.compile(
    r"(?P<code>`[^`\n]+`)|"
    r"(?P<strong>\*\*.+?\*\*|__.+?__)|"
    r"(?P<em>(?<!\*)\*[^*\n]+\*(?!\*)|(?<!_)_[^_\n]+_(?!_))"
)
_LIST = re.compile(r"^(?P<indent>\s*)(?P<marker>(?:\d+[.)]|[-+*]))\s+(?P<text>.*)$")
_FRONT_MATTER = re.compile(r"^---\s*$")


def _relative_source(path: Path) -> str:
    return path.resolve().relative_to(ROOT).as_posix()


def _slug(value: str) -> str:
    """Create a conservative, ASCII heading identifier."""

    value = re.sub(r"`([^`]*)`", r"\1", value)
    value = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", value)
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    value = re.sub(r"[^a-zA-Z0-9]+", "-", value).strip("-").lower()
    return value or "section"


def _page_prefix(path: Path) -> str:
    relative = path.resolve().relative_to(DOCS).with_suffix("")
    return "-".join(relative.parts)


def _prepare_page(path: Path) -> PreparedPage:
    text = path.read_text(encoding="utf-8")
    prefix = _page_prefix(path)
    counts: dict[str, int] = {}
    aliases: dict[str, list[str]] = {}
    rendered: list[str] = []
    fenced = False
    fence_mark = ""

    for line in text.splitlines():
        fence = _FENCE.match(line)
        if fence:
            marker = fence.group("mark")[0]
            if not fenced:
                fenced = True
                fence_mark = marker
            elif marker == fence_mark:
                fenced = False
            rendered.append(line)
            continue

        heading = None if fenced else _HEADING.match(line)
        if heading:
            heading_text = heading.group("text").strip()
            base = f"{prefix}-{_slug(heading_text)}"
            counts[base] = counts.get(base, 0) + 1
            anchor = base if counts[base] == 1 else f"{base}-{counts[base]}"
            aliases.setdefault(_slug(heading_text), []).append(anchor)
            aliases.setdefault(anchor, []).append(anchor)
            rendered.append(f"{heading.group('marks')} {heading_text} {{#{anchor}}}")
        else:
            rendered.append(line)

    if not aliases:
        # Every current source page has a heading, but retaining a deterministic
        # page anchor makes this helper safe for a future front-matter-only page.
        page_anchor = f"{prefix}-document"
    else:
        page_anchor = next(iter(aliases.values()))[0]

    return PreparedPage(
        path=path.resolve(),
        text="\n".join(rendered),
        prefix=prefix,
        page_anchor=page_anchor,
        anchors={key: tuple(value) for key, value in aliases.items()},
    )


def _anchor_for(page: PreparedPage, fragment: str) -> str | None:
    fragment = unquote(fragment).lstrip("#").casefold()
    if not fragment:
        return page.page_anchor
    exact = page.anchors.get(fragment)
    if exact and len(exact) == 1:
        return exact[0]

    # Existing source links may use an unprefixed auto-id, or a stable decision
    # ID such as ``global-01`` that prefixes the full heading slug.
    candidates = {
        anchor
        for alias, anchors in page.anchors.items()
        if alias.startswith(fragment + "-")
        for anchor in anchors
    }
    return next(iter(candidates)) if len(candidates) == 1 else None


def _resolve_target(page: PreparedPage, raw_path: str) -> Path | None:
    candidate = (page.path.parent / unquote(raw_path)).resolve()
    return candidate if candidate.exists() else None


def _generated_relative(path: Path) -> str:
    return quote(Path(os.path.relpath(path, OUTPUT)).as_posix(), safe="/")


def _rewrite_link(
    page: PreparedPage,
    target: str,
    pages: dict[Path, PreparedPage],
) -> str:
    target = target.strip()
    if target.startswith("<") and target.endswith(">"):
        target = target[1:-1]
    parsed = urlsplit(target)

    # Keep web, mailto, and other non-file links unchanged.
    if parsed.scheme or parsed.netloc:
        return target

    if not parsed.path:
        anchor = _anchor_for(page, parsed.fragment)
        return f"#{anchor}" if anchor is not None else target

    target_page = _resolve_target(page, parsed.path)
    if target_page is not None and target_page in pages:
        anchor = _anchor_for(pages[target_page], parsed.fragment)
        if anchor is not None:
            return f"#{anchor}"
        # An unresolved fragment is safer as a rebased file link than as a
        # guessed anchor; this keeps the source identity visible to the reader.
        if parsed.fragment:
            return f"{_generated_relative(target_page)}#{parsed.fragment}"
        return f"#{pages[target_page].page_anchor}"

    # Non-included repository resources (including README) are relative to the
    # generated output directory, not the temporary Markdown file or source doc.
    if target_page is None:
        target_page = (page.path.parent / unquote(parsed.path)).resolve()
    rebased = _generated_relative(target_page)
    return f"{rebased}#{parsed.fragment}" if parsed.fragment else rebased


def _rewrite_links(page: PreparedPage, pages: dict[Path, PreparedPage]) -> str:
    fenced = False
    fence_mark = ""
    lines: list[str] = []
    for line in page.text.splitlines():
        fence = _FENCE.match(line)
        if fence:
            marker = fence.group("mark")[0]
            if not fenced:
                fenced = True
                fence_mark = marker
            elif marker == fence_mark:
                fenced = False
            lines.append(line)
            continue
        if fenced:
            lines.append(line)
            continue

        def replace(match: re.Match[str]) -> str:
            new_target = _rewrite_link(page, match.group("target"), pages)
            return f"[{match.group('label')}]({new_target}{match.group('tail')})"

        lines.append(_LINK.sub(replace, line))
    return "\n".join(lines)


def combined_markdown(document: Document) -> str:
    pages = tuple(_prepare_page(page) for page in document.pages)
    page_map = {page.path: page for page in pages}
    sections = [
        "---",
        f'title: "{document.title}"',
        'subtitle: "January 2026"',
        "---",
        "",
    ]
    for page in pages:
        sections.extend(
            [
                "",
                '<div class="page-break"></div>',
                "",
                f"<!-- Included source: {_relative_source(page.path)} -->",
                _rewrite_links(page, page_map).strip(),
            ]
        )
    return "\n".join(sections) + "\n"


def _add_bookmark(paragraph, anchor: str, state: RenderState) -> None:
    name = state.bookmark_names.setdefault(
        anchor, "bm_" + re.sub(r"[^A-Za-z0-9_]", "_", anchor)
    )
    bookmark_id = str(state.next_bookmark_id)
    state.next_bookmark_id += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_hyperlink(paragraph, label: str, target: str, *, bold: bool = False) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    if target.startswith("#"):
        hyperlink.set(qn("w:anchor"), "bm_" + re.sub(r"[^A-Za-z0-9_]", "_", target[1:]))
    else:
        relationship_id = paragraph.part.relate_to(
            target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    if bold:
        properties.append(OxmlElement("w:b"))
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_inline(paragraph, text: str, *, bold: bool = False, italic: bool = False) -> None:
    """Add the small Markdown inline subset used by the runbooks."""

    position = 0
    while position < len(text):
        link = _LINK.search(text, position)
        marker = _INLINE_MARK.search(text, position)
        candidates = [match for match in (link, marker) if match is not None]
        if not candidates:
            plain = text[position:]
            run = paragraph.add_run(re.sub(r"\\([\\`*{}\[\]()#+.!_>-])", r"\1", plain))
            run.bold = bold
            run.italic = italic
            position = len(text)
            continue

        match = min(candidates, key=lambda item: item.start())
        if match.start() > position:
            plain = text[position : match.start()]
            run = paragraph.add_run(re.sub(r"\\([\\`*{}\[\]()#+.!_>-])", r"\1", plain))
            run.bold = bold
            run.italic = italic

        if match is link:
            target = match.group("target").strip()
            if target.startswith("<") and target.endswith(">"):
                target = target[1:-1]
            _add_hyperlink(paragraph, match.group("label"), target, bold=bold)
        else:
            value = match.group(0)
            if match.group("code") is not None:
                run = paragraph.add_run(value[1:-1])
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                run.bold = bold
                run.italic = italic
            elif match.group("strong") is not None:
                _add_inline(paragraph, value[2:-2], bold=True, italic=italic)
            else:
                _add_inline(paragraph, value[1:-1], bold=bold, italic=True)
        position = match.end()


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _split_table_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|") and not value.endswith("\\|"):
        value = value[:-1]
    return [cell.strip().replace("\\|", "|") for cell in value.split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _add_table(word_document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = word_document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            if column_index < len(row):
                _add_inline(paragraph, row[column_index])
            if row_index == 0:
                _set_cell_shading(cell, "D9EAF7")
                for run in paragraph.runs:
                    run.bold = True
    word_document.add_paragraph()


def _add_code_block(word_document, lines: list[str]) -> None:
    for line in lines:
        paragraph = word_document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        properties = run._r.get_or_add_rPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F2F2")
        properties.append(shading)


def _render_markdown(word_document, markdown: str) -> None:
    lines = markdown.splitlines()
    title = ""
    subtitle = ""
    index = 0
    if lines and _FRONT_MATTER.match(lines[0]):
        index = 1
        while index < len(lines) and not _FRONT_MATTER.match(lines[index]):
            key, separator, value = lines[index].partition(":")
            if separator:
                value = value.strip().strip('"')
                if key.strip() == "title":
                    title = value
                elif key.strip() == "subtitle":
                    subtitle = value
            index += 1
        index += 1

    if title:
        paragraph = word_document.add_paragraph(style="Title")
        _add_inline(paragraph, title)
    if subtitle:
        paragraph = word_document.add_paragraph(style="Subtitle")
        _add_inline(paragraph, subtitle)

    state = RenderState(bookmark_names={})
    while index < len(lines):
        line = lines[index]
        if not line.strip() or line.lstrip().startswith("<!--"):
            index += 1
            continue
        if line.strip() == '<div class="page-break"></div>':
            paragraph = word_document.add_paragraph()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            index += 1
            continue

        fence = _FENCE.match(line)
        if fence:
            marker = fence.group("mark")[0]
            indent = len(fence.group("indent"))
            index += 1
            code_lines: list[str] = []
            while index < len(lines):
                closing = _FENCE.match(lines[index])
                if closing and closing.group("mark")[0] == marker:
                    index += 1
                    break
                code_lines.append(lines[index][indent:] if lines[index].startswith(" " * indent) else lines[index])
                index += 1
            _add_code_block(word_document, code_lines)
            continue

        heading = _RENDER_HEADING.match(line)
        if heading:
            level = min(len(heading.group("marks")), 9)
            paragraph = word_document.add_paragraph(style=f"Heading {level}")
            paragraph.paragraph_format.keep_with_next = True
            _add_inline(paragraph, heading.group("text"))
            if heading.group("anchor"):
                _add_bookmark(paragraph, heading.group("anchor"), state)
            index += 1
            continue

        if index + 1 < len(lines) and "|" in line and _is_table_separator(lines[index + 1]):
            rows = [_split_table_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_split_table_row(lines[index]))
                index += 1
            _add_table(word_document, rows)
            continue

        list_item = _LIST.match(line)
        if list_item:
            marker = list_item.group("marker")
            content = [list_item.group("text")]
            index += 1
            while index < len(lines):
                continuation = lines[index]
                if not continuation.strip():
                    break
                if _LIST.match(continuation) or _FENCE.match(continuation) or _RENDER_HEADING.match(continuation):
                    break
                if len(continuation) - len(continuation.lstrip()) == 0:
                    break
                content.append(continuation.strip())
                index += 1
            paragraph = word_document.add_paragraph()
            prefix = "• " if marker in "-+*" else f"{marker} "
            if content[0].startswith("[ ] "):
                prefix += "☐ "
                content[0] = content[0][4:]
            elif content[0].startswith("[x] ") or content[0].startswith("[X] "):
                prefix += "☑ "
                content[0] = content[0][4:]
            _add_inline(paragraph, prefix + " ".join(content))
            continue

        if re.fullmatch(r"\s*(?:\*{3,}|-{3,}|_{3,})\s*", line):
            paragraph = word_document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run("―")
            index += 1
            continue

        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            if (
                not candidate.strip()
                or candidate.strip() == '<div class="page-break"></div>'
                or _FENCE.match(candidate)
                or _RENDER_HEADING.match(candidate)
                or _LIST.match(candidate)
                or (index + 1 < len(lines) and "|" in candidate and _is_table_separator(lines[index + 1]))
            ):
                break
            if candidate.lstrip().startswith("<!--"):
                index += 1
                continue
            paragraph_lines.append(candidate.strip())
            index += 1
        paragraph = word_document.add_paragraph()
        _add_inline(paragraph, " ".join(paragraph_lines))


def _configure_document(word_document) -> None:
    normal = word_document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    for section in word_document.sections:
        section.top_margin = Pt(54)
        section.bottom_margin = Pt(54)
        section.left_margin = Pt(54)
        section.right_margin = Pt(54)


def build() -> list[Path]:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    generated = []
    for document in DOCUMENTS:
        destination = OUTPUT / f"{document.stem}.docx"
        word_document = WordDocument()
        _configure_document(word_document)
        _render_markdown(word_document, combined_markdown(document))
        word_document.save(destination)
        generated.append(destination)
    return generated


def main() -> int:
    for path in build():
        print(path.relative_to(ROOT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
