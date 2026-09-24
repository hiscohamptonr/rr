"""Build the Excel-only BSCR population workbook.

Local maintainer command (not required on the reporting user's machine):
uv run --no-project --with openpyxl --with formulas --with numpy --with lxml \
    python reporting/bscr/tools/build_population_workbook.py

The delivered .xlsx uses ordinary Excel formulas and has no macros/external links.
"""

from __future__ import annotations

import argparse
import csv
import io
import math
import re
import zipfile
from copy import copy
from pathlib import Path

import formulas
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[3]
BSCR = ROOT / "reporting/bscr"
ENTITIES = ("33", "3624", "HIC", "HIG", "HSA")
# Only the visible schedules, not the vendor dropdown lists/XBRL helper grids below them.
SCHEDULE_ROWS = {
    "Schedule X(a)": 85,
    "Schedule X(b)": 85,
    "Schedule X(c)": 116,
    "Schedule X(f)": 80,
}
HEADERS = (
    "cntrycode",
    "bscr_entity",
    "region",
    "sum_pml",
    "sum_net",
    "count_policies",
    "is_geocoded",
    "peril_id",
)
AUTO = PatternFill("solid", fgColor="E2F0D9")
MANUAL = PatternFill("solid", fgColor="DDEBF7")
NAVY = PatternFill("solid", fgColor="17365D")
MONEY = "#,##0.000;[Red](#,##0.000);0.000"
XF_ROWS = {
    10: 1,
    11: 0,
    12: None,
    19: 1,
    20: 0,
    21: None,
    28: 1,
    33: 1,
    43: 0,
    46: 0,
    60: 0,
    64: 0,
}
REGIONS = [
    (27, "is_nahu", 2),
    (49, "is_na_eq", 1),
    (71, "is_eu", 2),
    (93, "is_jp_eq", 1),
    (115, "is_jp", 2),
]


def read_input(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if tuple(reader.fieldnames or ()) != HEADERS:
            raise ValueError(f"{path}: expected eight SQL output columns")
        rows = list(reader)
    seen = set()
    for r in rows:
        for k in ("sum_pml", "sum_net"):
            r[k] = float(r[k])
            if not math.isfinite(r[k]):
                raise ValueError(f"non-finite {k}")
        for k in ("count_policies", "is_geocoded", "peril_id"):
            r[k] = int(r[k])
        if r["is_geocoded"] not in (0, 1) or r["peril_id"] not in (1, 2):
            raise ValueError("invalid geocode/peril value")
        key = tuple(
            r[k]
            for k in ("cntrycode", "bscr_entity", "region", "is_geocoded", "peril_id")
        )
        if key in seen:
            raise ValueError(f"duplicate aggregate key: {key}")
        seen.add(key)
    if not rows:
        raise ValueError("empty SQL output")
    return rows


def sumifs(measure, entity, *, region="ALL", peril="1", geo=None, population="All"):
    criteria = f'BSCRInput[bscr_entity],"{entity}",BSCRInput[region],{region if region.startswith("IF(") else chr(34) + region + chr(34)},BSCRInput[peril_id],{peril}'
    if geo is not None:
        criteria += f",BSCRInput[is_geocoded],{geo}"
    total = f"SUMIFS(BSCRInput[{measure}],{criteria})"
    us = f'SUMIFS(BSCRInput[{measure}],{criteria},BSCRInput[cntrycode],"US")'
    expression = (
        us
        if population == "US"
        else f"({total}-{us})"
        if population == "Other"
        else total
    )
    return "=" + expression + ("" if measure == "count_policies" else "/1000000")


def auto(cell, formula, measure="sum_pml", note=""):
    cell.value = formula
    cell.fill = AUTO
    cell.protection = Protection(locked=True)
    cell.number_format = "#,##0" if measure == "count_policies" else MONEY
    cell.comment = Comment(
        note
        or "Automatic from SQL input. Money is already USD: divide by 1,000,000 only. Counts are unscaled.",
        "BSCR mapping",
    )


def title(ws, value, end=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end)
    c = ws.cell(1, 1, value)
    c.fill = NAVY
    c.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws.row_dimensions[1].height = 32
    ws.sheet_view.showGridLines = False


def create_workbook(rows):
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(
        calcId=191029, fullCalcOnLoad=True, forceFullCalc=True
    )
    info = wb.create_sheet("Instructions")
    title(info, "BSCR — automatic schedule population")
    info["A4"] = "Reporting year (1 January as-at)"
    info["B4"] = 2026
    info["A5"] = "USD to USD millions"
    info["B5"] = 1000000
    info["A6"] = "Input basis"
    info["B6"] = "Current SQL"
    info["B4"].fill = MANUAL
    info["B6"].fill = MANUAL
    mode = DataValidation(type="list", formula1='"Current SQL,Legacy check"')
    info.add_data_validation(mode)
    mode.add(info["B6"])
    instructions = [
        "1. SQL input is preloaded with the supplied current SQL output. To refresh, clear ALL old table data rows, then paste the eight SQL columns as VALUES below the headers. Keep the Excel table and its headers.",
        "2. Use the full output for all five entities (@bscr_entity = NULL). The Excel table expands when rows are pasted. Do not leave old rows beneath a shorter replacement.",
        "3. Current SQL mode: region/peril routing matches bscr-output.sql. ALL uses earthquake. Legacy check mode is only for bscr-extract-legacy.sql output already converted to USD at FX 1.35.",
        "4. Green cells are automatic and locked. Blue cells require manual input: premiums, model/EP losses, questionnaires and narratives. Blank manual cells do NOT mean zero. There are no carried-forward historical premium or loss values.",
        "5. Agreed classification: geocoded = modellable/modelled/detailed; ungeocoded = not modellable/not modelled/data deficient. All other contracts = ALL minus US. Counts use count_policies, not independently verified distinct contracts.",
        "6. Amounts in the SQL input are FULL USD. Formulas divide by 1,000,000 once; they NEVER apply another FX conversion. SQL exposure is placed in the historical all-other-lines fields. Statutory property-cat fields need separate manual data.",
        "7. Each entity tab stacks X(a), X(b), X(c) and X(f). Repeated schedule populations are intentional; do not sum regions or schedule sections together.",
        "8. Historical checks compare saved April figures, the historical baseline and the LIVE generated cells. Classification differences are expected where April was manually reallocated. Corrected-input differences are not hidden.",
        "9. Recalculate in Excel after pasting (Automatic calculation / Calculate Now). This is a working population workbook, not a complete or approved regulatory return. Review manual fields before transfer.",
    ]
    for row, text in enumerate(instructions, 9):
        info.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        info.cell(row, 1, text).alignment = Alignment(wrap_text=True, vertical="center")
        info.row_dimensions[row].height = 46
    info["A20"] = "Input rows"
    info["B20"] = "=COUNTA(BSCRInput[region])"
    info["A21"] = "Unmapped/unrecognized entity rows"
    info["B21"] = "=COUNTA(BSCRInput[region])-" + "-".join(
        f'COUNTIF(BSCRInput[bscr_entity],"{e}")' for e in ENTITIES
    )
    info["A22"] = "Blank-entity gross USDm (overlapping regions)"
    info["B22"] = '=SUMIFS(BSCRInput[sum_pml],BSCRInput[bscr_entity],"")/1000000'
    info["A23"] = "Missing geocode flags"
    info["B23"] = "=COUNTBLANK(BSCRInput[is_geocoded])"
    info["A24"] = "Invalid peril values"
    info["B24"] = (
        "=COUNTA(BSCRInput[region])-COUNTIF(BSCRInput[peril_id],1)-COUNTIF(BSCRInput[peril_id],2)"
    )
    for i, ent in enumerate(ENTITIES, 26):
        info.cell(i, 1, ent + " ALL rows")
        info.cell(
            i,
            2,
            f'COUNTIFS(BSCRInput[bscr_entity],"{ent}",BSCRInput[region],"ALL",BSCRInput[peril_id],1)',
        )
        info.cell(i, 2).value = "=" + info.cell(i, 2).value
    info["A32"] = (
        "Do not ignore unmapped exposure. Missing entity ALL rows indicate incomplete input. Blank-entity exposure is not assigned to any entity tab."
    )
    info.merge_cells("A32:H33")
    info["A32"].alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 48
    info.column_dimensions["B"].width = 26
    for col in "CDEFGH":
        info.column_dimensions[col].width = 13
    info.merge_cells("D20:H24")
    info["D20"] = (
        '=IF(OR(B21>0,B23>0,B24>0,COUNTIF(B26:B30,0)>0,SUM(B34:B37)>0),"REVIEW INPUT: unmapped rows, invalid fields or missing entities. See controls at left.","INPUT CONTROLS PASS: still complete and review all blue manual fields.")'
    )
    info["D20"].alignment = Alignment(wrap_text=True, vertical="center")
    warning = PatternFill("solid", fgColor="FFC7CE")
    info.conditional_formatting.add(
        "B21 B23:B24", CellIsRule(operator="greaterThan", formula=["0"], fill=warning)
    )
    info.conditional_formatting.add(
        "B26:B30", CellIsRule(operator="equal", formula=["0"], fill=warning)
    )
    for row, field, label in [
        (34, "sum_pml", "Missing / nonnumeric gross amounts"),
        (35, "sum_net", "Missing / nonnumeric net amounts"),
        (36, "count_policies", "Missing / nonnumeric counts"),
    ]:
        info.cell(row, 1, label)
        info.cell(row, 2, f"=COUNTA(BSCRInput[region])-COUNT(BSCRInput[{field}])")
    info["A37"] = "Invalid geocode flags"
    info["B37"] = (
        "=COUNTA(BSCRInput[region])-COUNTIF(BSCRInput[is_geocoded],0)-COUNTIF(BSCRInput[is_geocoded],1)"
    )
    info.conditional_formatting.add(
        "B34:B37", CellIsRule(operator="greaterThan", formula=["0"], fill=warning)
    )
    info.freeze_panes = "C9"

    inp = wb.create_sheet("SQL input")
    inp.append(HEADERS)
    for r in rows:
        inp.append([r[k] for k in HEADERS])
    table = Table(displayName="BSCRInput", ref=f"A1:H{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    inp.add_table(table)
    inp.freeze_panes = "D2"
    for col in "ABCDEFGH":
        inp.column_dimensions[col].width = 23
    inp["D1"].comment = Comment(
        "Full USD, not USD millions. Already converted by SQL.", "Input contract"
    )
    inp["F1"].comment = Comment(
        "Existing grouped-source-row count; not a distinct-contract measure.",
        "Input contract",
    )
    for row in inp.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0" if cell.column == 6 else "#,##0.00"
    for column, choices in [("G", '"0,1"'), ("H", '"1,2"')]:
        validation = DataValidation(type="list", formula1=choices, allow_blank=False)
        validation.errorTitle = "Invalid SQL flag"
        validation.error = "Paste the unchanged SQL output flags."
        validation.showErrorMessage = True
        inp.add_data_validation(validation)
        validation.add(f"{column}2:{column}1048576")
    source_maps = {}
    auto_specs = []
    for ent in ENTITIES:
        source = load_workbook(
            BSCR / "reconcile" / f"2026 BSCR - UKEU - {ent}.xlsx", data_only=False
        )
        cached = load_workbook(
            BSCR / "reconcile" / f"2026 BSCR - UKEU - {ent}.xlsx", data_only=True
        )
        ws = wb.create_sheet(ent)
        title(ws, f"{ent} — BSCR working schedules", 13)
        ws["A2"] = "Entity"
        ws["B2"] = ent
        ws["D2"] = "Input basis"
        ws["E2"] = "='Instructions'!B6"
        ws["G2"] = "Coverage"
        ws["H2"] = (
            f'=IF(COUNTIFS(BSCRInput[bscr_entity],"{ent}",BSCRInput[region],"ALL",BSCRInput[peril_id],1)=0,"NO INPUT FOR ENTITY","INPUT PRESENT")'
        )
        ws.merge_cells("A3:M4")
        ws["A3"] = (
            "Green = automatic. Blue = manual, not supplied by SQL. Counts/geocoding use the agreed historical proxy; premiums and EP losses are NOT automatically complete."
        )
        ws["A3"].alignment = Alignment(wrap_text=True)
        for col in range(1, 14):
            ws.column_dimensions[chr(64 + col)].width = 16
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 5
        ws.column_dimensions["C"].width = 13
        ws.column_dimensions["D"].width = 36
        ws.column_dimensions["I"].width = 35
        offset = 7
        for sh in ("Schedule X(a)", "Schedule X(b)", "Schedule X(c)", "Schedule X(f)"):
            src = source[sh]
            old = cached[sh]
            source_maps[(ent, sh)] = offset
            end_row = SCHEDULE_ROWS[sh]
            for row in src.iter_rows(max_row=end_row, max_col=13):
                for c in row:
                    if c.__class__.__name__ == "MergedCell":
                        continue
                    dest = ws.cell(c.row + offset, c.column)
                    dest.font = copy(c.font)
                    dest.fill = copy(c.fill)
                    dest.border = copy(c.border)
                    dest.number_format = c.number_format
                    dest.alignment = copy(c.alignment)
                    dest.protection = copy(c.protection)
                    if not c.protection.locked:
                        dest.value = None
                        dest.fill = MANUAL
                        dest.comment = Comment(
                            f"Manual input: not supplied by SQL. Historical {sh}!{c.coordinate} cached value: {old[c.coordinate].value!s}. Do not assume it remains appropriate.",
                            "Manual field",
                        )
                    elif c.data_type == "f":
                        if (
                            "[" in c.value
                            or "#REF!" in c.value
                            or "HYPERLINK" in c.value
                        ):
                            dest.value = None
                        else:
                            dest.value = Translator(
                                c.value, origin=c.coordinate
                            ).translate_formula(dest.coordinate)
                    elif c.value != "HICM":
                        dest.value = c.value
            for merged in src.merged_cells.ranges:
                if merged.max_row > end_row or merged.max_col > 13:
                    continue
                ws.merge_cells(
                    start_row=merged.min_row + offset,
                    start_column=merged.min_col,
                    end_row=merged.max_row + offset,
                    end_column=merged.max_col,
                )
            for r, dim in src.row_dimensions.items():
                if r <= end_row and dim.height:
                    ws.row_dimensions[r + offset].height = dim.height
            ws.cell(2 + offset, 2, ent)
            ws.cell(3 + offset, 2, "=DATE('Instructions'!B4,1,1)")
            if sh in ("Schedule X(c)", "Schedule X(f)"):
                c = ws.cell(3 + offset, 3)
                if c.__class__.__name__ != "MergedCell":
                    c.value = None
            link = ws.cell(1 + offset, 7 if sh == "Schedule X(c)" else 10)
            if link.__class__.__name__ != "MergedCell":
                link.value = "Back to instructions"
                link.hyperlink = "#'Instructions'!A1"

            def put(addr, formula, measure="sum_pml", spec=None):
                cell = src[addr]
                target = ws.cell(cell.row + offset, cell.column)
                auto(target, formula, measure)
                if spec is not None:
                    auto_specs.append(
                        {
                            "sheet": ent,
                            "cell": target.coordinate,
                            "source_sheet": sh,
                            "source_cell": addr,
                            "measure": measure,
                            **spec,
                        }
                    )

            if sh in ("Schedule X(a)", "Schedule X(b)"):
                for row, geo in (
                    [(25, 1), (26, 0)] if sh == "Schedule X(a)" else [(27, 1), (28, 0)]
                ):
                    for col, m in [("E", "sum_pml"), ("J", "sum_net")]:
                        put(
                            f"{col}{row}",
                            sumifs(m, ent, geo=geo),
                            m,
                            {
                                "region": "ALL",
                                "peril": 1,
                                "geo": geo,
                                "population": "All",
                            },
                        )
            if sh == "Schedule X(c)":
                for row, region, peril in REGIONS:
                    reg = (
                        'IF(Instructions!$B$6="Legacy check","is_jp","is_jp_eq")'
                        if region == "is_jp_eq"
                        else region
                    )
                    per = (
                        'IF(Instructions!$B$6="Legacy check",1,2)'
                        if peril == 2
                        else "1"
                    )
                    for col, m in [("E", "sum_pml"), ("F", "sum_net")]:
                        put(
                            f"{col}{row}",
                            sumifs(m, ent, region=reg, peril=per),
                            m,
                            {
                                "region": region,
                                "peril": peril,
                                "geo": None,
                                "population": "All",
                            },
                        )
            if sh == "Schedule X(f)":
                for row, geo in XF_ROWS.items():
                    for pop, cols in [("US", "EFG"), ("Other", "HIJ"), ("All", "KLM")]:
                        for col, m in zip(
                            cols, ("count_policies", "sum_pml", "sum_net")
                        ):
                            put(
                                f"{col}{row}",
                                sumifs(m, ent, geo=geo, population=pop),
                                m,
                                {
                                    "region": "ALL",
                                    "peril": 1,
                                    "geo": geo,
                                    "population": pop,
                                },
                            )
                # Historical convention places these populations wholly in detailed/data-deficient.
                for row in (29, 30, 31, 32, 44, 45, 61, 62, 63):
                    for col in "EFGHIJKLM":
                        cell = ws.cell(row + offset, ord(col) - 64)
                        if cell.__class__.__name__ != "MergedCell":
                            auto(
                                cell,
                                "=0",
                                "count_policies" if col in "EHK" else "sum_pml",
                                "Zero under the agreed geocode-proxy allocation; does not classify April manual reallocations.",
                            )
            navcol = {
                "Schedule X(a)": 2,
                "Schedule X(b)": 5,
                "Schedule X(c)": 8,
                "Schedule X(f)": 11,
            }[sh]
            ws.cell(6, navcol, sh).hyperlink = f"#'{ent}'!A{offset + 1}"
            ws.row_breaks.append(Break(id=offset + end_row))
            offset += end_row + 5
        ws.freeze_panes = "E8"
        ws.sheet_view.zoomScale = 65
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = f"A1:M{offset}"
        ws.protection.sheet = True
        source.close()
        cached.close()
    return wb, source_maps, auto_specs


class Evaluator:
    """Evaluate the workbook's real formula text using an independent Excel-function engine."""

    def __init__(self, wb, rows):
        self.wb = wb
        self.values = {}
        self.compiled = {}
        self.sheets = {s.upper(): s for s in wb.sheetnames}
        self.arrays = {
            "INPUT_" + k.upper(): np.array([r[k] for r in rows], dtype=object).reshape(
                -1, 1
            )
            for k in HEADERS
        }

    def cell(self, sheet, addr):
        key = (sheet, addr.replace("$", ""))
        if key in self.values:
            return self.values[key]
        value = self.wb[sheet][key[1]].value
        if isinstance(value, str) and value.startswith("="):
            text = re.sub(
                r"BSCRInput\[([^\]]+)\]",
                lambda m: "INPUT_" + m.group(1).upper(),
                value,
                flags=re.I,
            )
            if text not in self.compiled:
                self.compiled[text] = formulas.Parser().ast(text)[1].compile()
            fn = self.compiled[text]
            result = fn(*[self.resolve(sheet, name) for name in fn.inputs])
            value = np.asarray(result).item()
            if isinstance(value, np.generic):
                value = value.item()
            if isinstance(value, str) and value.startswith(
                ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!")
            ):
                raise ValueError((sheet, addr, value))
        if value is None:
            value = ""
        self.values[key] = value
        return value

    def resolve(self, sheet, name):
        if name in self.arrays:
            return self.arrays[name]
        name = name.replace("$", "")
        if "!" in name:
            sh, name = name.rsplit("!", 1)
            sheet = self.sheets[sh.strip("'").upper()]
        if ":" in name:
            from openpyxl.utils.cell import range_boundaries

            a, b, c, d = range_boundaries(name)
            return np.array(
                [
                    [
                        self.cell(sheet, self.wb[sheet].cell(r, col).coordinate)
                        for col in range(a, c + 1)
                    ]
                    for r in range(b, d + 1)
                ],
                dtype=object,
            )
        return self.cell(sheet, name)


def add_historical_checks(wb, maps):
    ws = wb.create_sheet("Historical checks")
    title(ws, "Historical checks — expected differences are visible", 9)
    ws.merge_cells("A2:I3")
    ws["A2"] = (
        "April templates are the layout/reference. Legacy baseline uses old saved CSV x1.35. This tab does not force current corrected output to equal history. Classification splits use the agreed geocode proxy and can differ from April manual reallocations."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    headers = [
        "Entity",
        "Field",
        "Measure",
        "April saved",
        "Legacy baseline",
        "April minus legacy",
        "Live workbook",
        "Live minus legacy",
        "Interpretation",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(5, i, h).fill = NAVY
        ws.cell(5, i).font = Font(bold=True, color="FFFFFF")
    old = list(csv.DictReader((BSCR / "output/bscr-output.csv").open()))
    targets = []
    for ent in ENTITIES:
        source = load_workbook(
            BSCR / "reconcile" / f"2026 BSCR - UKEU - {ent}.xlsx", data_only=True
        )
        for row, geo in [(12, None), (10, 1), (11, 0)]:
            for col, m in [("K", "count_policies"), ("L", "sum_pml"), ("M", "sum_net")]:
                targets.append(
                    (
                        ent,
                        "Schedule X(f)",
                        f"{col}{row}",
                        m,
                        "ALL",
                        geo,
                        source["Schedule X(f)"][f"{col}{row}"].value,
                        "ALL total" if row == 12 else "Geocode-proxy classification",
                    )
                )
        for row, region, peril in REGIONS:
            for col, m in [("E", "sum_pml"), ("F", "sum_net")]:
                targets.append(
                    (
                        ent,
                        "Schedule X(c)",
                        f"{col}{row}",
                        m,
                        "is_jp" if region == "is_jp_eq" else region,
                        None,
                        source["Schedule X(c)"][f"{col}{row}"].value,
                        region,
                    )
                )
        source.close()
    for row, (ent, sh, addr, m, region, geo, saved, label) in enumerate(targets, 6):
        baseline = sum(
            float(r[m])
            for r in old
            if r["bscr_entity"] == ent
            and r["region"] == region
            and (geo is None or int(r["is_geocoded"]) == geo)
        )
        if m != "count_policies":
            baseline *= 1.35 / 1000000
        target = (
            wb[ent]
            .cell(
                int(re.search(r"\d+", addr).group()) + maps[(ent, sh)],
                ord(addr[0]) - 64,
            )
            .coordinate
        )
        values = [
            ent,
            f"{sh}!{addr} — {label}",
            m,
            saved,
            baseline,
            f"=D{row}-E{row}" if saved is not None else None,
            f"='{ent}'!{target}",
            f"=G{row}-E{row}",
            "April classification may differ"
            if "classification" in label
            else "Rounded April total"
            if label == "ALL total" and m != "count_policies"
            else "Compare region/peril basis",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        for col in range(4, 9):
            ws.cell(row, col).number_format = (
                "#,##0" if m == "count_policies" else MONEY
            )
        ws.row_dimensions[row].height = 32
        for col in (2, 9):
            ws.cell(row, col).alignment = Alignment(wrap_text=True)
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 22
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["I"].width = 35
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:I{ws.max_row}"
    ws.sheet_view.showGridLines = False
    ws.conditional_formatting.add(
        f"H6:H{ws.max_row}",
        FormulaRule(
            formula=['ABS(H6)>IF($C6="count_policies",0,0.00000001)'],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )
    ws.conditional_formatting.add(
        f"H6:H{ws.max_row}",
        FormulaRule(
            formula=['ABS(H6)<=IF($C6="count_policies",0,0.00000001)'], fill=AUTO
        ),
    )
    return targets


def verify(wb, rows, maps, specs):
    ev = Evaluator(wb, rows)
    for spec in specs:
        region = spec["region"]
        peril = spec["peril"]
        if wb["Instructions"]["B6"].value == "Legacy check":
            peril = 1
            if region == "is_jp_eq":
                region = "is_jp"
        expected = sum(
            r[spec["measure"]]
            for r in rows
            if r["bscr_entity"] == spec["sheet"]
            and r["region"] == region
            and r["peril_id"] == peril
            and (spec["geo"] is None or r["is_geocoded"] == spec["geo"])
            and (
                spec["population"] == "All"
                or (r["cntrycode"] == "US") == (spec["population"] == "US")
            )
        )
        if spec["measure"] != "count_policies":
            expected /= 1000000
        actual = ev.cell(spec["sheet"], spec["cell"])
        tolerance = 0 if spec["measure"] == "count_policies" else 0.00000001
        if abs(float(actual) - expected) > tolerance:
            raise AssertionError((spec, actual, expected))
    for ws in wb:
        for row in ws:
            for c in row:
                if c.data_type == "f":
                    ev.cell(ws.title, c.coordinate)
    return ev


def save_cached(wb, ev, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    stream = io.BytesIO()
    wb.save(stream)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    n = "{" + ns["s"] + "}"
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as z:
        entries = {i.filename: (i, z.read(i.filename)) for i in z.infolist()}
    for i, ws in enumerate(wb, 1):
        key = f"xl/worksheets/sheet{i}.xml"
        # Preserve namespaces/prefixes on workbook XML using lxml in-place below.
        from lxml import etree as L

        root = L.fromstring(entries[key][1])
        for c in root.findall(".//s:sheetData/s:row/s:c", ns):
            if c.find("s:f", ns) is None:
                continue
            value = ev.cell(ws.title, c.attrib["r"])
            v = c.find("s:v", ns)
            if v is None:
                v = L.SubElement(c, n + "v")
            if isinstance(value, str):
                c.set("t", "str")
                v.text = value
            elif isinstance(value, bool):
                c.set("t", "b")
                v.text = "1" if value else "0"
            else:
                c.attrib.pop("t", None)
                v.text = repr(float(value))
        entries[key] = (
            entries[key][0],
            L.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
        )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for info, data in entries.values():
            z.writestr(info, data)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input", type=Path, default=BSCR / "reconcile/bscr_output_wseq.csv"
    )
    parser.add_argument(
        "--output", type=Path, default=BSCR / "workbooks/BSCR_Auto_Population.xlsx"
    )
    args = parser.parse_args()
    rows = read_input(args.input)
    wb, maps, specs = create_workbook(rows)
    add_historical_checks(wb, maps)
    current = verify(wb, rows, maps, specs)
    # Independent historical data exercise: the same Excel formulas, not a second mapping implementation.
    legacy = []
    with (BSCR / "output/bscr-output.csv").open() as f:
        for r in csv.DictReader(f):
            for k in ("sum_pml", "sum_net"):
                r[k] = float(r[k]) * 1.35
            for k in ("count_policies", "is_geocoded"):
                r[k] = int(r[k])
            r["peril_id"] = 1
            legacy.append(r)
    wb["Instructions"]["B6"] = "Legacy check"
    historical = verify(wb, legacy, maps, specs)
    checks = wb["Historical checks"]
    for row in range(6, checks.max_row + 1):
        if abs(float(historical.cell(checks.title, f"H{row}"))) > 0.00000001:
            raise AssertionError(("legacy mismatch", row))
    wb["Instructions"]["B6"] = "Current SQL"
    current = verify(wb, rows, maps, specs)
    save_cached(wb, current, args.output)
    saved = load_workbook(args.output, data_only=True)
    assert saved.sheetnames == [
        "Instructions",
        "SQL input",
        *ENTITIES,
        "Historical checks",
    ]
    for ent in ENTITIES:
        addr = saved[ent].cell(12 + maps[(ent, "Schedule X(f)")], 11).coordinate
        assert saved[ent][addr].value == sum(
            r["count_policies"]
            for r in rows
            if r["bscr_entity"] == ent and r["region"] == "ALL" and r["peril_id"] == 1
        )
    with zipfile.ZipFile(args.output) as z:
        assert z.testzip() is None
        assert not any("externalLinks/" in p for p in z.namelist())
    print(
        f"Created {args.output}: 5 entity sheets, {len(specs)} verified automatic cells; current and legacy formula scenarios passed."
    )


if __name__ == "__main__":
    main()
